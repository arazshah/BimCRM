import datetime
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.club.models import PointEntry
from apps.core.tenancy import tenant_context
from apps.crm.models import Customer, Policy
from apps.finance.models import CommissionRule, Installment
from apps.panel.importer import clean_mobile, clean_national_code, parse_amount, parse_date

from .conftest import host
from .test_addons import enable

pytestmark = pytest.mark.django_db

CUST = ["موبایل", "نام و نام خانوادگی", "کد ملی", "شهر"]
POL = ["شماره بیمه‌نامه", "موبایل بیمه‌گذار", "نام بیمه‌گذار", "رشته", "شرکت بیمه", "تاریخ شروع", "حق بیمه (ریال)", "پلاک"]
INST = ["شماره بیمه‌نامه", "تاریخ سررسید", "مبلغ (ریال)", "وضعیت", "شماره چک", "بانک"]


def xlsx(**sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("data.xlsx", buf.getvalue())


def upload(client, tenant, f):
    return client.post("/panel/import/", {"action": "preview", "file": f}, HTTP_HOST=host(tenant))


def confirm(client, tenant):
    return client.post("/panel/import/", {"action": "confirm"}, HTTP_HOST=host(tenant))


def full_workbook():
    return xlsx(
        **{
            "مشتریان": [CUST, ["9121112233", "سارا محمدی", 12345678, "تهران"], ["09121112244", "علی", "", ""]],
            "بیمه‌نامه‌ها": [
                POL,
                ["1403-55", "09121112233", "", "ثالث", "ایران", "1403/05/12", "45,000,000", "۱۲ ب ۳۴۵"],
                ["1403-66", "۰۹۱۲۵۵۵۶۶۷۷", "رضا نیک", "بدنه", "بیمه آسیا", "14030601", 12000000, ""],
            ],
            "اقساط": [
                INST,
                ["1403-55", "1403/06/12", 15000000, "پرداخت شده", "", ""],
                ["1403-55", "1403/07/12", 15000000, "", "778899", "ملت"],
            ],
        }
    )


def test_parsers():
    assert clean_mobile(9121234567) == "09121234567"
    assert clean_mobile("+98 912 123 4567") == "09121234567"
    assert clean_national_code(12345678) == "0012345678"
    assert parse_amount("۴۵٬۰۰۰٬۰۰۰ ریال") == 45_000_000
    assert parse_date("1403/05/12") == datetime.date(2024, 8, 2)
    assert parse_date("14030512") == parse_date("12/05/1403") == parse_date("۱۴۰۳-۵-۱۲")
    assert parse_date(datetime.datetime(2024, 8, 2)) == datetime.date(2024, 8, 2)
    with pytest.raises(ValueError):
        parse_date("1403/13/40")


def test_template_download_has_all_sheets(client, tenant, owner):
    enable(tenant, "finance")
    client.force_login(owner)
    resp = client.get("/panel/import/template.xlsx", HTTP_HOST=host(tenant))
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["راهنما", "مشتریان", "بیمه‌نامه‌ها", "اقساط", "مقادیر مجاز"]
    assert [c.value for c in wb["مشتریان"][1]][:2] == ["موبایل", "نام و نام خانوادگی"]
    page = client.get("/panel/import/", HTTP_HOST=host(tenant)).content.decode()
    assert "دریافت قالب اکسل" in page and "نگران تکراری شدن نباشید" in page


def test_template_without_finance_has_no_installments(client, tenant, owner):
    client.force_login(owner)
    wb = load_workbook(io.BytesIO(client.get("/panel/import/template.xlsx", HTTP_HOST=host(tenant)).content))
    assert "اقساط" not in wb.sheetnames


def test_preview_writes_nothing_then_confirm_imports_everything(client, tenant, owner):
    enable(tenant, "finance")
    with tenant_context(tenant):
        CommissionRule.objects.create(rate=Decimal("12.5"))
    client.force_login(owner)
    page = upload(client, tenant, full_workbook()).content.decode()
    assert "پیش‌نمایش" in page and "ثبت نهایی" in page
    with tenant_context(tenant):
        assert not Customer.objects.exists() and not Policy.objects.exists()
    confirm(client, tenant)
    with tenant_context(tenant):
        sara = Customer.objects.get(mobile="09121112233")
        assert sara.full_name == "سارا محمدی" and sara.national_code == "0012345678"
        assert sara.source == Customer.Source.IMPORT
        reza = Customer.objects.get(mobile="09125556677")
        assert reza.full_name == "رضا نیک"
        assert Customer.objects.count() == 3
        p = Policy.objects.get(number="1403-55")
        assert p.customer == sara and p.line.slug == "third-party" and p.company.name == "بیمه ایران"
        assert p.premium == 45_000_000 and p.start_date == datetime.date(2024, 8, 2)
        assert p.end_date == datetime.date(2025, 8, 3)  # one Jalali year later: 1404/05/12
        assert p.created_by == owner and p.subject == "۱۲ ب ۳۴۵"
        assert p.commission_rate == Decimal("12.5")
        assert Policy.objects.get(number="1403-66").company.name == "بیمه آسیا"
        items = list(Installment.objects.filter(policy=p).order_by("due_date"))
        assert [i.status for i in items] == ["paid", "pending"]
        assert items[0].paid_at == items[0].due_date
        assert items[1].method == "cheque" and items[1].cheque_number == "778899" and items[1].cheque_bank == "ملت"
        assert not PointEntry.objects.exists()  # history does not earn club points


def test_reimport_never_duplicates_and_updates(client, tenant, owner):
    enable(tenant, "finance")
    client.force_login(owner)
    upload(client, tenant, full_workbook())
    confirm(client, tenant)
    page = upload(client, tenant, full_workbook()).content.decode()
    assert 'value="confirm"' not in page  # nothing changed
    changed = xlsx(**{"بیمه‌نامه‌ها": [POL, ["1403-55", "09121112233", "", "ثالث", "ایران", "1403/05/12", 50000000, ""]]})
    upload(client, tenant, changed)
    confirm(client, tenant)
    with tenant_context(tenant):
        assert Customer.objects.count() == 3 and Policy.objects.count() == 2 and Installment.objects.count() == 2
        p = Policy.objects.get(number="1403-55")
        assert p.premium == 50_000_000 and p.subject == "۱۲ ب ۳۴۵"  # empty cell keeps old value


def test_fixes_question_mark_names_from_bad_csv(client, tenant, owner):
    client.force_login(owner)
    with tenant_context(tenant):
        Customer.objects.create(full_name="???? ??????", mobile="09121112233")
        Customer.objects.create(full_name="??? ????", mobile="09121112244", city="کرج")
    f = xlsx(
        Sheet1=[["نام", "نام خانوادگی", "موبایل"], ["سارا", "محمدی", "09121112233"], ["علی", "رضایی", "09121112244"]]
    )
    page = upload(client, tenant, f).content.decode()
    assert "به‌روز می‌شود" in page
    confirm(client, tenant)
    with tenant_context(tenant):
        assert Customer.objects.count() == 2
        assert dict(Customer.objects.values_list("mobile", "full_name")) == {
            "09121112233": "سارا محمدی",
            "09121112244": "علی رضایی",
        }
        assert Customer.objects.get(mobile="09121112244").city == "کرج"


def test_garbled_names_are_rejected_not_saved(client, tenant, owner):
    client.force_login(owner)
    data = "نام و نام خانوادگی,موبایل\n???? ??????,09121112233\n".encode("cp1256", errors="replace")
    page = upload(client, tenant, SimpleUploadedFile("c.csv", data)).content.decode()
    assert "علامت سؤال" in page
    confirm(client, tenant)
    with tenant_context(tenant):
        assert not Customer.objects.exists()


def test_csv_utf8_and_windows_encodings(client, tenant, owner):
    client.force_login(owner)
    # Windows "CSV (Comma delimited)" on a Persian-locale PC: cp1256 with Arabic ي/ك
    rows = "نام و نام خانوادگی;موبایل;شهر\nمینا کریمی;09121231234;تهران\n".replace("ی", "ي").replace("ک", "ك")
    upload(client, tenant, SimpleUploadedFile("a.csv", rows.encode("cp1256")))
    confirm(client, tenant)
    rows2 = "نام,موبایل,کد ملی,شهر\nمینا,09121231235,0012345678,تهران\nتکراری,09121231235,,\nبد,123,,\n"
    page = upload(client, tenant, SimpleUploadedFile("b.csv", rows2.encode("utf-8-sig"))).content.decode()
    assert "تکرار شده" in page and "موبایل نامعتبر" in page
    confirm(client, tenant)
    with tenant_context(tenant):
        assert dict(Customer.objects.values_list("mobile", "full_name")) == {
            "09121231234": "مینا کریمی",
            "09121231235": "مینا",
        }


def test_missing_required_column_and_bad_rows_are_reported(client, tenant, owner):
    enable(tenant, "finance")
    client.force_login(owner)
    f = xlsx(
        **{
            "بیمه‌نامه‌ها": [
                ["شماره بیمه‌نامه", "موبایل بیمه‌گذار", "رشته", "شرکت بیمه", "تاریخ شروع", "ستون اضافه"],
                ["A1", "09121112233", "فضایی", "ایران", "1403/01/01", "x"],
                ["A2", "09121112233", "ثالث", "ناشناس", "1403/01/01", "x"],
                ["A3", "09121112233", "ثالث", "ایران", "1403/15/01", "x"],
                ["A4", "09121112233", "ثالث", "ایران", "1403/01/01", "x"],
            ],
            "اقساط": [["شماره بیمه‌نامه", "مبلغ (ریال)"], ["A4", 100]],
        }
    )
    page = upload(client, tenant, f).content.decode()
    assert "رشته «فضایی» شناخته نشد" in page and "شرکت بیمه «ناشناس» شناخته نشد" in page
    assert "تاریخ شروع نامعتبر" in page and "ستون اضافه" in page
    assert "ستون‌های اجباری پیدا نشد: تاریخ سررسید" in page
    confirm(client, tenant)
    with tenant_context(tenant):
        assert list(Policy.objects.values_list("number", flat=True)) == ["A4"]
        assert Customer.objects.get().full_name == "09121112233"


def test_installments_need_finance(client, tenant, owner):
    client.force_login(owner)
    page = upload(client, tenant, full_workbook()).content.decode()
    assert "افزونه مالی" in page
    confirm(client, tenant)
    with tenant_context(tenant):
        assert Policy.objects.count() == 2 and not Installment.objects.exists()


def test_import_is_isolated_per_office(client, tenant, other_tenant, owner):
    with tenant_context(other_tenant):
        Customer.objects.create(full_name="دیگری", mobile="09121112233")
    client.force_login(owner)
    upload(client, tenant, xlsx(Sheet1=[CUST, ["09121112233", "سارا", "", ""]]))
    confirm(client, tenant)
    with tenant_context(other_tenant):
        assert Customer.objects.get().full_name == "دیگری"
    with tenant_context(tenant):
        assert Customer.objects.get().full_name == "سارا"


def test_wrong_file_type_and_cancel(client, tenant, owner):
    client.force_login(owner)
    page = upload(client, tenant, SimpleUploadedFile("a.pdf", b"%PDF")).content.decode()
    assert "فقط فایل اکسل" in page
    upload(client, tenant, xlsx(Sheet1=[CUST, ["09121112233", "سارا", "", ""]]))
    client.post("/panel/import/", {"action": "cancel"}, HTTP_HOST=host(tenant))
    confirm(client, tenant)
    with tenant_context(tenant):
        assert not Customer.objects.exists()


def test_filled_template_round_trip(client, tenant, owner):
    enable(tenant, "finance")
    client.force_login(owner)
    wb = load_workbook(io.BytesIO(client.get("/panel/import/template.xlsx", HTTP_HOST=host(tenant)).content))
    ws = wb["بیمه‌نامه‌ها"]
    titles = [c.value for c in ws[1]]
    row = {"شماره بیمه‌نامه": "T-1", "موبایل بیمه‌گذار": "09121112233", "نام بیمه‌گذار": "سارا"}
    row |= {"رشته": "شخص ثالث", "شرکت بیمه": "البرز", "تاریخ شروع": "1403/05/12"}
    for i, t in enumerate(titles, start=1):
        ws.cell(2, i).value = row.get(t)
    buf = io.BytesIO()
    wb.save(buf)
    upload(client, tenant, SimpleUploadedFile("filled.xlsx", buf.getvalue()))
    confirm(client, tenant)
    with tenant_context(tenant):
        p = Policy.objects.get()
        assert p.number == "T-1" and p.customer.full_name == "سارا" and p.company.name == "بیمه البرز"
