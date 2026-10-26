"""Bulk import from one standard Excel template: customers, policies and installments.

The office fills our template (fixed column titles), uploads it, sees a preview, then confirms.
Rows are matched, never duplicated: customers by mobile, policies by number + company,
installments by policy + due date. Re-uploading the same file only applies changes.
"""

import contextlib
import csv
import datetime
import io
import logging
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import jdatetime

from apps.accounts.models import normalize_mobile
from apps.core.models import InsuranceCompany, LineOfBusiness

logger = logging.getLogger(__name__)

MAX_ROWS = 10_000
MAX_BYTES = 10 * 1024 * 1024

_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
_ARABIC = str.maketrans({"ي": "ی", "ى": "ی", "ك": "ک", "ة": "ه", "ۀ": "ه", "‏": "", "‎": ""})


class ImportFileError(Exception):
    """The whole file can't be read (wrong type, broken encoding, no known sheet)."""


# ---------------------------------------------------------------- template schema


@dataclass(frozen=True)
class Col:
    key: str
    title: str
    required: bool = False
    help: str = ""
    example: str = ""
    aliases: tuple = ()


@dataclass(frozen=True)
class Sheet:
    key: str
    title: str
    intro: str
    columns: tuple
    feature: str = ""

    def required(self):
        return [c for c in self.columns if c.required]


CUSTOMERS = Sheet(
    "customers",
    "مشتریان",
    "هر ردیف یک مشتری. موبایل کلید شناسایی است؛ اگر مشتری از قبل باشد، اطلاعاتش به‌روز می‌شود و تکراری ساخته نمی‌شود.",
    (
        Col(
            "mobile",
            "موبایل",
            True,
            "۱۱ رقم؛ با یا بدون صفر اول، فارسی یا انگلیسی",
            "09121234567",
            ("شماره موبایل", "تلفن همراه", "همراه", "شماره همراه", "موبایل بیمه گذار"),
        ),
        Col(
            "full_name",
            "نام و نام خانوادگی",
            True,
            "یا به‌جای آن دو ستون «نام» و «نام خانوادگی» بگذارید",
            "علی رضایی",
            ("نام کامل", "نام بیمه گذار"),
        ),
        Col("first_name", "نام", False, "فقط اگر نام و نام خانوادگی را جدا دارید", "علی"),
        Col("last_name", "نام خانوادگی", False, "فقط اگر نام و نام خانوادگی را جدا دارید", "رضایی", ("فامیل",)),
        Col(
            "national_code",
            "کد ملی",
            False,
            "۱۰ رقم؛ صفرهای اولِ حذف‌شده خودکار برمی‌گردد",
            "0012345678",
            ("کدملی", "شماره ملی", "کد ملی بیمه گذار"),
        ),
        Col("birth_date", "تاریخ تولد", False, "شمسی، مثل ۱۳۶۵/۰۴/۲۰", "1365/04/20"),
        Col("city", "شهر", False, "", "تهران"),
        Col("notes", "یادداشت", False, "هر توضیحی درباره مشتری", "مشتری قدیمی", ("توضیحات", "توضیح")),
    ),
)

POLICIES = Sheet(
    "policies",
    "بیمه‌نامه‌ها",
    "هر ردیف یک بیمه‌نامه. اگر بیمه‌گذار در برگه مشتریان نباشد، با همین موبایل و نام ساخته می‌شود. "
    "بیمه‌نامه با شماره و شرکت شناسایی می‌شود و ورود دوباره، آن را به‌روز می‌کند.",
    (
        Col(
            "number",
            "شماره بیمه‌نامه",
            True,
            "همان شماره روی بیمه‌نامه",
            "1403/123/4567",
            ("شماره بیمه نامه", "ش بیمه نامه", "کد یکتا", "شماره"),
        ),
        Col(
            "mobile",
            "موبایل بیمه‌گذار",
            True,
            "برای پیدا کردن یا ساختن مشتری",
            "09121234567",
            ("موبایل", "شماره موبایل", "تلفن همراه", "همراه"),
        ),
        Col(
            "full_name",
            "نام بیمه‌گذار",
            False,
            "برای مشتری جدید لازم است؛ اگر نباشد موبایل جای نام می‌نشیند",
            "علی رضایی",
            ("نام و نام خانوادگی", "بیمه گذار", "نام کامل"),
        ),
        Col("national_code", "کد ملی بیمه‌گذار", False, "", "0012345678", ("کد ملی", "کدملی")),
        Col(
            "line",
            "رشته",
            True,
            "مثل ثالث، بدنه، آتش‌سوزی، عمر، درمان، مسافرتی، مسئولیت",
            "ثالث",
            ("رشته بیمه", "نوع بیمه", "نوع بیمه نامه"),
        ),
        Col(
            "company", "شرکت بیمه", True, "مثل ایران، آسیا، البرز (کلمه «بیمه» لازم نیست)", "ایران", ("شرکت", "بیمه گر")
        ),
        Col("start_date", "تاریخ شروع", True, "شمسی، مثل ۱۴۰۳/۰۵/۱۲", "1403/05/12", ("تاریخ صدور", "شروع", "از تاریخ")),
        Col(
            "end_date",
            "تاریخ پایان",
            False,
            "اگر خالی باشد، یک سال بعد از شروع",
            "1404/05/12",
            ("تاریخ انقضا", "سررسید بیمه نامه", "پایان", "تا تاریخ"),
        ),
        Col(
            "premium",
            "حق بیمه (ریال)",
            False,
            "عدد به ریال؛ ویرگول و جداکننده اشکالی ندارد",
            "45000000",
            ("حق بیمه", "مبلغ حق بیمه", "مبلغ کل"),
        ),
        Col(
            "subject",
            "موضوع بیمه / پلاک",
            False,
            "پلاک خودرو، نشانی ملک و…",
            "۱۲ ب ۳۴۵ ایران ۶۷",
            ("پلاک", "شماره پلاک", "موضوع", "موضوع بیمه", "مورد بیمه"),
        ),
        Col("status", "وضعیت", False, "فعال یا ابطال؛ خالی یعنی فعال", "فعال"),
        Col(
            "commission_rate",
            "نرخ کارمزد (درصد)",
            False,
            "خالی = از جدول نرخ کارمزد",
            "10",
            ("نرخ کارمزد", "درصد کارمزد"),
        ),
    ),
)

INSTALLMENTS = Sheet(
    "installments",
    "اقساط",
    "هر ردیف یک قسط یا چک. بیمه‌نامه با شماره‌اش پیدا می‌شود (در همین فایل یا از قبل در سیستم). "
    "قسط با بیمه‌نامه و تاریخ سررسید شناسایی می‌شود و ورود دوباره، آن را به‌روز می‌کند.",
    (
        Col(
            "number",
            "شماره بیمه‌نامه",
            True,
            "همان شماره در برگه بیمه‌نامه‌ها",
            "1403/123/4567",
            ("شماره بیمه نامه", "شماره"),
        ),
        Col("company", "شرکت بیمه", False, "فقط اگر یک شماره در دو شرکت تکرار شده", "ایران", ("شرکت",)),
        Col("due_date", "تاریخ سررسید", True, "شمسی", "1403/06/12", ("سررسید", "تاریخ قسط", "تاریخ چک")),
        Col("amount", "مبلغ (ریال)", True, "عدد به ریال", "15000000", ("مبلغ", "مبلغ قسط", "مبلغ چک")),
        Col("status", "وضعیت", False, "در انتظار، پرداخت‌شده یا برگشتی؛ خالی یعنی در انتظار", "در انتظار"),
        Col("paid_at", "تاریخ پرداخت", False, "برای قسط پرداخت‌شده", "1403/06/10", ("تاریخ وصول",)),
        Col("method", "روش پرداخت", False, "نقد، کارت‌خوان، واریز، درگاه یا چک", "چک", ("روش", "نحوه پرداخت")),
        Col("cheque_number", "شماره چک", False, "", "123456", ("سریال چک",)),
        Col("cheque_bank", "بانک", False, "", "ملت", ("نام بانک",)),
        Col("note", "توضیح", False, "", "", ("توضیحات", "یادداشت")),
    ),
    feature="finance",
)

SHEETS = (CUSTOMERS, POLICIES, INSTALLMENTS)


# ---------------------------------------------------------------- cleaning helpers


def _key(text) -> str:
    """Comparable form of a column title or a name: Persian letters, no spaces, no (…)/*."""
    s = str(text or "").translate(_ARABIC).translate(_DIGITS)
    s = re.sub(r"\(.*?\)", "", s)
    s = re.sub(r"[\s‌*:_\-/.]+", "", s)
    return s.lower()


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).translate(_ARABIC)
    return re.sub(r"\s+", " ", s).strip()


def is_garbled(text: str) -> bool:
    """Names a wrong CSV encoding turned into ????."""
    return bool(text) and ("?" in text or "�" in text) and not re.search(r"[A-Za-z؀-ۿ]", text)


def _digits(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value).translate(_DIGITS))


def clean_mobile(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return normalize_mobile(str(value or "").translate(_DIGITS).replace(".", ""))


def valid_mobile(mobile: str) -> bool:
    return len(mobile) == 11 and mobile.startswith("09") and mobile.isdigit()


def clean_national_code(value) -> str:
    d = _digits(value)
    if not d:
        return ""
    if len(d) < 10:
        d = d.zfill(10)
    return d[:10]


def parse_amount(value) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).translate(_DIGITS)
    s = re.sub(r"[,\s٬،ریالتومان]", "", s)
    if not s:
        return None
    try:
        return int(Decimal(s))
    except InvalidOperation:
        raise ValueError("عدد نیست") from None


def parse_date(value) -> datetime.date | None:
    """Jalali text (1403/05/12, 1403-5-12, 14030512, 12/05/1403) or a real Excel date."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime.datetime):
        value = value.date()
    if isinstance(value, datetime.date):
        if value.year < 1700:  # Jalali digits typed into a date cell
            return jdatetime.date(value.year, value.month, value.day).togregorian()
        return value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).translate(_DIGITS).strip().split(" ")[0]
    parts = [p for p in re.split(r"[/\-.\\]", s) if p]
    if len(parts) == 1 and len(parts[0]) == 8:
        parts = [parts[0][:4], parts[0][4:6], parts[0][6:]]
    if len(parts) != 3 or not all(p.isdigit() for p in parts):
        raise ValueError("تاریخ نامعتبر")
    a, b, c = (int(p) for p in parts)
    if c > 1000 and a <= 31:  # day/month/year
        a, c = c, a
    if a < 100:
        a += 1300 if a > 50 else 1400
    try:
        if 1200 <= a < 1700:
            return jdatetime.date(a, b, c).togregorian()
        if 1900 <= a < 2200:
            return datetime.date(a, b, c)
    except ValueError:
        pass
    raise ValueError("تاریخ نامعتبر")


def _choice(value: str, table: dict, label: str):
    k = _key(value)
    if not k:
        return None
    for word, code in table.items():
        if _key(word) in k:
            return code
    raise ValueError(f"{label} «{value}» شناخته نشد")


# ---------------------------------------------------------------- reading files


def read_rows(uploaded_name: str, data: bytes) -> list[tuple[str, list[list]]]:
    """[(sheet name, rows)] from .xlsx or .csv. The first non-empty row is the header."""
    name = (uploaded_name or "").lower()
    if len(data) > MAX_BYTES:
        raise ImportFileError("حجم فایل بیش از ۱۰ مگابایت است؛ آن را در چند فایل جدا وارد کنید.")
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return _read_xlsx(data)
    if name.endswith(".xls"):
        raise ImportFileError("فرمت قدیمی xls پشتیبانی نمی‌شود. در اکسل Save As › Excel Workbook (.xlsx) بزنید.")
    if name.endswith((".csv", ".txt")):
        return [("", _read_csv(data))]
    raise ImportFileError("فقط فایل اکسل (.xlsx) یا CSV پذیرفته می‌شود.")


def _read_xlsx(data: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any broken workbook
        raise ImportFileError("فایل اکسل باز نشد. مطمئن شوید فایل خراب نیست و با فرمت .xlsx ذخیره شده.") from exc
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS + 50:
                break
        out.append((ws.title, rows))
    wb.close()
    return out


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp1256"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("کدگذاری فایل CSV شناخته نشد. فایل را به‌صورت اکسل (.xlsx) بفرستید.")


def _read_csv(data: bytes):
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


# ---------------------------------------------------------------- import


@dataclass
class SheetReport:
    sheet: Sheet
    found: bool = False
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list = field(default_factory=list)  # (row number, message)
    ignored_columns: list = field(default_factory=list)
    skipped_reason: str = ""

    @property
    def total(self):
        return self.created + self.updated + self.unchanged + len(self.errors)

    def error(self, row, msg):
        self.errors.append((row, msg))


@dataclass
class Report:
    sheets: list

    @property
    def any_found(self):
        return any(s.found for s in self.sheets)

    @property
    def changes(self):
        return sum(s.created + s.updated for s in self.sheets)

    @property
    def error_count(self):
        return sum(len(s.errors) for s in self.sheets)


def _header_map(sheet: Sheet, header: list) -> tuple[dict, list]:
    lookup = {}
    for col in sheet.columns:
        for name in (col.title, *col.aliases):
            lookup.setdefault(_key(name), col.key)
    mapping, ignored = {}, []
    for idx, title in enumerate(header):
        k = _key(title)
        if not k:
            continue
        key = lookup.get(k)
        if key and key not in mapping:
            mapping[key] = idx
        else:
            ignored.append(clean_text(title))
    return mapping, ignored


def _guess_sheet(name: str, header: list) -> Sheet | None:
    k = _key(name)
    if "قسط" in k or "اقساط" in k or "چک" in k:
        return INSTALLMENTS
    if "بیمهنامه" in k:
        return POLICIES
    if "مشتری" in k or "مشتریان" in k:
        return CUSTOMERS
    titles = {_key(h) for h in header}
    if "تاریخسررسید" in titles or "سررسید" in titles:
        return INSTALLMENTS
    if {"شمارهبیمهنامه", "رشته"} & titles:
        return POLICIES
    if "موبایل" in titles:
        return CUSTOMERS
    return None


def _split(rows):
    """Header row = first row that has at least two non-empty cells."""
    for i, row in enumerate(rows):
        if sum(1 for c in row if c not in (None, "")) >= 2:
            return row, rows[i + 1 :], i + 2
    return [], [], 0


class Importer:
    def __init__(self, tenant, *, user=None, branch=None, features=(), can_finance=True):
        self.tenant = tenant
        self.user = user
        self.branch = branch
        self.features = set(features)
        self.can_finance = can_finance
        self.commit = False
        self._customers = None
        self._policies = None
        self._lines = None
        self._companies = None
        self._rates = {}
        # New rows are inserted in batches at the end of each sheet: thousands of rows stay well under the
        # request timeout. Updates to existing rows are saved one by one.
        self._pending = {"customers": [], "policies": [], "installments": []}

    # ----- lookups (loaded once; new rows are added as they are created)

    def _load(self):
        from apps.crm.models import Customer, Policy

        self._customers = {c.mobile: c for c in Customer.objects.all()}
        self._policies = {}
        for p in Policy.objects.select_related("company").order_by("created_at"):
            self._policies[(p.number, p.company_id)] = p
            self._policies[(p.number, None)] = p
        self._lines = list(LineOfBusiness.objects.all())
        self._companies = list(InsuranceCompany.objects.all())

    def match_line(self, text):
        k = _key(text)
        if not k:
            raise ValueError("رشته خالی است")
        for line in self._lines:
            if k in (_key(line.name), _key(line.slug)):
                return line
        hits = [ln for ln in self._lines if k in _key(ln.name) or _key(ln.name) in k]
        if len(hits) == 1:
            return hits[0]
        raise ValueError(f"رشته «{text}» شناخته نشد")

    def match_company(self, text):
        def bare(s):
            return _key(s).replace("شرکت", "").replace("بیمه", "")

        k = bare(text)
        if not k:
            raise ValueError("شرکت بیمه خالی است")
        for c in self._companies:
            if k == bare(c.name):
                return c
        hits = [c for c in self._companies if k in bare(c.name) or bare(c.name) in k]
        if len(hits) == 1:
            return hits[0]
        raise ValueError(f"شرکت بیمه «{text}» شناخته نشد")

    # ----- entry point

    def run(self, sheets_rows, *, commit: bool) -> Report:
        from django.db import transaction

        self.commit = commit
        self._load()
        reports = {s.key: SheetReport(s) for s in SHEETS}
        parsed = []
        for name, rows in sheets_rows:
            header, body, first_row = _split(rows)
            if not header:
                continue
            sheet = _guess_sheet(name, header)
            if sheet is None or reports[sheet.key].found:
                continue
            reports[sheet.key].found = True
            parsed.append((sheet, header, body, first_row))
        if not parsed:
            raise ImportFileError(
                "هیچ برگه‌ای با ستون‌های قالب پیدا نشد. قالب آماده را از همین صفحه بگیرید و عنوان ستون‌ها را تغییر ندهید."
            )
        order = {s.key: i for i, s in enumerate(SHEETS)}
        parsed.sort(key=lambda p: order[p[0].key])
        with transaction.atomic():
            for sheet, header, body, first_row in parsed:
                self._run_sheet(reports[sheet.key], header, body, first_row)
            if not commit:
                transaction.set_rollback(True)
        return Report([reports[s.key] for s in SHEETS])

    def _run_sheet(self, rep: SheetReport, header, body, first_row):
        sheet = rep.sheet
        if sheet.feature and (sheet.feature not in self.features or not self.can_finance):
            rep.skipped_reason = "افزونه مالی برای این دفتر فعال نیست یا اجازه ورود اقساط را ندارید."
            return
        mapping, rep.ignored_columns = _header_map(sheet, header)
        missing = [
            c.title
            for c in sheet.required()
            if c.key not in mapping and not (c.key == "full_name" and "first_name" in mapping)
        ]
        if missing:
            rep.skipped_reason = "ستون‌های اجباری پیدا نشد: " + "، ".join(missing)
            return
        handler = getattr(self, f"_row_{sheet.key}")
        data_rows = [(first_row + i, r) for i, r in enumerate(body) if any(c not in (None, "") for c in r)]
        if len(data_rows) > MAX_ROWS:
            rep.skipped_reason = f"بیش از {MAX_ROWS:,} ردیف دارد؛ فایل را چند تکه کنید."
            return
        seen = set()
        self._rows(rep, handler, mapping, data_rows, seen)
        self._flush()

    def _flush(self):
        from apps.crm.models import Customer, Policy
        from apps.finance.models import Installment

        if not self.commit:
            return
        for key, model in (("customers", Customer), ("policies", Policy), ("installments", Installment)):
            objs = self._pending[key]
            if objs:
                model.objects.bulk_create(objs, batch_size=500)
                objs.clear()

    def _rows(self, rep, handler, mapping, data_rows, seen):
        from django.db import DatabaseError, transaction

        for rownum, raw in data_rows:
            row = {k: (raw[i] if i < len(raw) else None) for k, i in mapping.items()}
            try:
                # Savepoint only when writing; the preview never touches the database.
                with transaction.atomic() if self.commit else contextlib.nullcontext():
                    result = handler(row, seen)
            except ValueError as exc:
                rep.error(rownum, str(exc))
                continue
            except DatabaseError:
                logger.exception("import row %s failed", rownum)
                rep.error(rownum, "ذخیره این ردیف ممکن نشد؛ مقادیر را بررسی کنید")
                continue
            setattr(rep, result, getattr(rep, result) + 1)

    # ----- rows

    def _customer(self, mobile, *, full_name="", national_code="", fields=None):
        """Create or update the customer for this mobile. Returns (customer, 'created'|'updated'|'unchanged')."""
        from apps.crm.models import Customer

        fields = dict(fields or {})
        if full_name:
            fields["full_name"] = full_name[:120]
        if national_code:
            fields["national_code"] = national_code
        customer = self._customers.get(mobile)
        if customer is None:
            customer = Customer(
                tenant=self.tenant,
                mobile=mobile,
                full_name=fields.pop("full_name", "") or mobile,
                source=Customer.Source.IMPORT,
                branch=self.branch,
                **fields,
            )
            if self.commit:
                self._pending["customers"].append(customer)
            self._customers[mobile] = customer
            return customer, "created"
        changed = [k for k, v in fields.items() if v not in ("", None) and getattr(customer, k) != v]
        for k in changed:
            setattr(customer, k, fields[k])
        if changed and self.commit and customer.pk:
            customer.save(update_fields=changed)
        return customer, ("updated" if changed else "unchanged")

    def _name(self, value, required_msg=None):
        name = clean_text(value)
        if is_garbled(name):
            raise ValueError("نام به‌صورت علامت سؤال است (فایل با کدگذاری اشتباه ذخیره شده)")
        if not name and required_msg:
            raise ValueError(required_msg)
        return name

    def _row_customers(self, row, seen):
        mobile = clean_mobile(row.get("mobile"))
        if not valid_mobile(mobile):
            raise ValueError(f"موبایل نامعتبر: {clean_text(row.get('mobile')) or 'خالی'}")
        if mobile in seen:
            raise ValueError("این موبایل در همین فایل تکرار شده؛ ردیف اول ثبت شد")
        seen.add(mobile)
        name = self._name(row.get("full_name"))
        if not name:
            name = " ".join(p for p in (self._name(row.get("first_name")), self._name(row.get("last_name"))) if p)
        if not name and mobile not in self._customers:
            raise ValueError("نام خالی است")
        fields = {}
        city = clean_text(row.get("city"))[:60]
        if city:
            fields["city"] = city
        notes = clean_text(row.get("notes"))
        if notes:
            fields["notes"] = notes
        if row.get("birth_date") not in (None, ""):
            try:
                fields["birth_date"] = parse_date(row.get("birth_date"))
            except ValueError:
                raise ValueError("تاریخ تولد نامعتبر") from None
        _, result = self._customer(
            mobile, full_name=name, national_code=clean_national_code(row.get("national_code")), fields=fields
        )
        return result

    def _row_policies(self, row, seen):
        from apps.crm.models import Policy
        from apps.finance.models import CommissionRule

        number = clean_text(row.get("number"))[:60]
        if not number:
            raise ValueError("شماره بیمه‌نامه خالی است")
        mobile = clean_mobile(row.get("mobile"))
        if not valid_mobile(mobile):
            raise ValueError(f"موبایل بیمه‌گذار نامعتبر: {clean_text(row.get('mobile')) or 'خالی'}")
        line = self.match_line(row.get("line"))
        company = self.match_company(row.get("company"))
        if (number, company.pk) in seen:
            raise ValueError("این بیمه‌نامه در همین فایل تکرار شده؛ ردیف اول ثبت شد")
        seen.add((number, company.pk))
        try:
            start = parse_date(row.get("start_date"))
        except ValueError:
            raise ValueError("تاریخ شروع نامعتبر") from None
        if not start:
            raise ValueError("تاریخ شروع خالی است")
        try:
            end = parse_date(row.get("end_date"))
        except ValueError:
            raise ValueError("تاریخ پایان نامعتبر") from None
        if not end:
            from apps.core import jalali

            end = jalali.add_months(start, 12)
        if end < start:
            raise ValueError("تاریخ پایان قبل از شروع است")
        try:
            premium = parse_amount(row.get("premium"))
        except ValueError:
            raise ValueError("حق بیمه عدد نیست") from None
        if premium is not None and premium < 0:
            raise ValueError("حق بیمه منفی است")
        status = _choice(
            clean_text(row.get("status")),
            {"ابطال": "cancelled", "باطل": "cancelled", "غیرفعال": "cancelled", "فعال": "active"},
            "وضعیت",
        )
        rate = None
        if row.get("commission_rate") not in (None, ""):
            try:
                rate = Decimal(str(row.get("commission_rate")).translate(_DIGITS).replace("%", "").strip())
            except InvalidOperation:
                raise ValueError("نرخ کارمزد عدد نیست") from None
            if not (0 <= rate <= 100):
                raise ValueError("نرخ کارمزد باید بین ۰ و ۱۰۰ باشد")
            rate = rate.quantize(Decimal("0.01"))
        customer, _ = self._customer(
            mobile,
            full_name=self._name(row.get("full_name")),
            national_code=clean_national_code(row.get("national_code")),
        )
        values = {
            "line": line,
            "start_date": start,
            "end_date": end,
            "subject": clean_text(row.get("subject"))[:160],
        }
        if premium is not None:
            values["premium"] = premium
        if status:
            values["status"] = status
        if rate is not None:
            values["commission_rate"] = rate
        policy = self._policies.get((number, company.pk))
        if policy is None:
            policy = Policy(
                tenant=self.tenant,
                customer=customer,
                company=company,
                number=number,
                branch=self.branch,
                created_by=self.user,
                **values,
            )
            # Batched insert: no club points or SMS for history; fill what Policy.save() would.
            if policy.commission_rate is None:
                pair = (company.pk, line.pk)
                if pair not in self._rates:
                    self._rates[pair] = CommissionRule.rate_for(self.tenant.pk, company.pk, line.pk)
                policy.commission_rate = self._rates[pair]
            policy.marketer_id = customer.marketer_id
            policy.branch_id = self.branch.pk if self.branch else customer.branch_id
            if self.commit:
                self._pending["policies"].append(policy)
            self._policies[(number, company.pk)] = policy
            self._policies.setdefault((number, None), policy)
            return "created"
        if not values["subject"]:
            values.pop("subject")
        if customer.pk and policy.customer_id != customer.pk:
            values["customer"] = customer
        changed = [k for k, v in values.items() if getattr(policy, k) != v]
        for k in changed:
            setattr(policy, k, values[k])
        if changed and self.commit and policy.pk:
            policy.save(update_fields=changed)
        return "updated" if changed else "unchanged"

    def _row_installments(self, row, seen):
        from apps.finance.models import Installment

        number = clean_text(row.get("number"))
        if not number:
            raise ValueError("شماره بیمه‌نامه خالی است")
        company_id = None
        if clean_text(row.get("company")):
            company_id = self.match_company(row.get("company")).pk
        policy = self._policies.get((number, company_id))
        if policy is None:
            raise ValueError(f"بیمه‌نامه {number} نه در فایل هست نه در سیستم")
        try:
            due = parse_date(row.get("due_date"))
        except ValueError:
            raise ValueError("تاریخ سررسید نامعتبر") from None
        if not due:
            raise ValueError("تاریخ سررسید خالی است")
        try:
            amount = parse_amount(row.get("amount"))
        except ValueError:
            raise ValueError("مبلغ عدد نیست") from None
        if not amount or amount <= 0:
            raise ValueError("مبلغ خالی یا صفر است")
        status = (
            _choice(
                clean_text(row.get("status")),
                {
                    "برگشت": Installment.Status.BOUNCED,
                    "انتظار": Installment.Status.PENDING,
                    "نشده": Installment.Status.PENDING,
                    "باز": Installment.Status.PENDING,
                    "پرداخت": Installment.Status.PAID,
                    "وصول": Installment.Status.PAID,
                    "تسویه": Installment.Status.PAID,
                },
                "وضعیت",
            )
            or Installment.Status.PENDING
        )
        method = _choice(
            clean_text(row.get("method")),
            {
                "چک": Installment.Method.CHEQUE,
                "نقد": Installment.Method.CASH,
                "کارتخوان": Installment.Method.CARD,
                "پوز": Installment.Method.CARD,
                "کارت": Installment.Method.CARD,
                "واریز": Installment.Method.TRANSFER,
                "حواله": Installment.Method.TRANSFER,
                "درگاه": Installment.Method.ONLINE,
                "آنلاین": Installment.Method.ONLINE,
            },
            "روش پرداخت",
        )
        cheque_number = clean_text(row.get("cheque_number"))[:30]
        if cheque_number and not method:
            method = Installment.Method.CHEQUE
        try:
            paid_at = parse_date(row.get("paid_at"))
        except ValueError:
            raise ValueError("تاریخ پرداخت نامعتبر") from None
        if status == Installment.Status.PAID and not paid_at:
            paid_at = due
        key = (id(policy), due, cheque_number)
        if key in seen:
            raise ValueError("این قسط در همین فایل تکرار شده؛ ردیف اول ثبت شد")
        seen.add(key)
        values = {"amount": amount, "status": status, "paid_at": paid_at if status == "paid" else None}
        if method:
            values["method"] = method
        if cheque_number:
            values["cheque_number"] = cheque_number
        bank = clean_text(row.get("cheque_bank"))[:60]
        if bank:
            values["cheque_bank"] = bank
        note = clean_text(row.get("note"))[:200]
        if note:
            values["note"] = note
        existing = None
        if policy.pk:
            qs = Installment.objects.filter(policy=policy, due_date=due)
            if cheque_number:
                qs = qs.filter(cheque_number__in=[cheque_number, ""])
            existing = qs.order_by("-cheque_number", "id").first()
        if existing is None:
            if self.commit:
                self._pending["installments"].append(
                    Installment(tenant=self.tenant, policy=policy, due_date=due, **values)
                )
            return "created"
        changed = [k for k, v in values.items() if getattr(existing, k) != v]
        for k in changed:
            setattr(existing, k, values[k])
        if changed and self.commit:
            existing.save(update_fields=changed)
        return "updated" if changed else "unchanged"


# ---------------------------------------------------------------- template workbook


def build_template(include_finance: bool = True) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    guide = wb.active
    guide.title = "راهنما"
    guide.sheet_view.rightToLeft = True
    bold = Font(bold=True)
    req_fill = PatternFill("solid", fgColor="F8D7DA")
    opt_fill = PatternFill("solid", fgColor="E2EFDA")
    head_font = Font(bold=True, size=12)
    wrap = Alignment(wrap_text=True, vertical="top")

    lines = [
        ("راهنمای ورود اطلاعات از اکسل", None),
        ("", None),
        (
            "۱. اطلاعات را در برگه‌های «مشتریان»، «بیمه‌نامه‌ها» و «اقساط» بنویسید. هر برگه‌ای که ندارید را خالی بگذارید.",
            None,
        ),
        ("۲. عنوان ستون‌ها را تغییر ندهید. ترتیب ستون‌ها مهم نیست و ستون اضافه نادیده گرفته می‌شود.", None),
        ("۳. ستون‌های قرمز اجباری‌اند و ستون‌های سبز اختیاری؛ ستون اختیاری‌ای که ندارید را می‌توانید حذف کنید.", None),
        (
            "۴. نگران تکراری نباشید: مشتری با «موبایل»، بیمه‌نامه با «شماره بیمه‌نامه + شرکت» و قسط با «بیمه‌نامه + تاریخ سررسید» "
            "شناسایی می‌شود. اگر از قبل باشد به‌روز می‌شود و دوباره ساخته نمی‌شود.",
            None,
        ),
        ("۵. می‌توانید فقط برگه بیمه‌نامه‌ها را پر کنید؛ بیمه‌گذار با موبایل و نامش خودکار ساخته می‌شود.", None),
        ("۶. تاریخ‌ها شمسی (مثل 1403/05/12) و مبالغ به ریال. اعداد فارسی یا انگلیسی فرقی نمی‌کند.", None),
        (
            "۷. فایل را با همین فرمت .xlsx ذخیره کنید و در پنل › ورود از اکسل بارگذاری کنید. قبل از ثبت، پیش‌نمایش نشان داده می‌شود.",
            None,
        ),
        ("", None),
    ]
    for text, _ in lines:
        guide.append([text])
    guide["A1"].font = Font(bold=True, size=14)
    sheets = [s for s in SHEETS if include_finance or not s.feature]
    for sheet in sheets:
        guide.append([f"برگه «{sheet.title}»"])
        guide.cell(guide.max_row, 1).font = head_font
        guide.append([sheet.intro])
        guide.append(["عنوان ستون", "اجباری؟", "توضیح", "نمونه"])
        for c in guide[guide.max_row]:
            c.font = bold
        for col in sheet.columns:
            guide.append([col.title, "اجباری" if col.required else "اختیاری", col.help, col.example])
            guide.cell(guide.max_row, 1).fill = req_fill if col.required else opt_fill
        guide.append([])
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 70
    guide.column_dimensions["D"].width = 22
    for row in guide.iter_rows():
        for c in row:
            c.alignment = wrap

    for sheet in sheets:
        ws = wb.create_sheet(sheet.title)
        ws.sheet_view.rightToLeft = True
        ws.append([c.title for c in sheet.columns])
        for i, col in enumerate(sheet.columns, start=1):
            cell = ws.cell(1, i)
            cell.font = bold
            cell.fill = req_fill if col.required else opt_fill
            note = ("اجباری. " if col.required else "اختیاری. ") + col.help
            if col.example:
                note += f"\nنمونه: {col.example}"
            cell.comment = Comment(note, "BimCRM")
            letter = cell.column_letter
            ws.column_dimensions[letter].width = max(14, len(col.title) + 4)
            # Keep mobiles, national codes, numbers and dates as text so Excel doesn't drop leading zeros.
            for r in range(2, 2001):
                ws.cell(r, i).number_format = "@"
        ws.freeze_panes = "A2"

    ref = wb.create_sheet("مقادیر مجاز")
    ref.sheet_view.rightToLeft = True
    ref.append(["رشته‌ها", "شرکت‌های بیمه"])
    for c in ref[1]:
        c.font = bold
    lines_ = list(LineOfBusiness.objects.filter(is_active=True).values_list("name", flat=True))
    companies = list(InsuranceCompany.objects.filter(is_active=True).values_list("name", flat=True))
    for i in range(max(len(lines_), len(companies))):
        ref.append([lines_[i] if i < len(lines_) else None, companies[i] if i < len(companies) else None])
    ref.column_dimensions["A"].width = 24
    ref.column_dimensions["B"].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
